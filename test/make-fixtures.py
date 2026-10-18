# 全部為虛構資料
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
import os; FX = os.path.join(os.path.dirname(__file__), "fixtures") + "/"; os.makedirs(FX, exist_ok=True)
def roster(fn, prefix, names):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "姓名查詢"
    ws.append(["姓名", "建制隊號", "女生隊號", "隊別", "學籍號碼"])
    for i, n in enumerate(names, 1):
        ws.append([n, int(f"{prefix}{i:04d}"), i, 1, int(f"1{prefix}99{i:05d}")])
    wb.save(FX + fn)
roster("roster44.xlsx", 44, ["王小美", "林小華", "陳小玉", "同名測試"])
roster("roster45.xlsx", 45, ["張小芳", "李小婷", "黃小琪", "同名測試"])

# 加扣分範本
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "工作表1"
hdr = ["學籍號碼", "姓名", "加減分依據", "事蹟說明", "說明", "加扣分", "分數"]
ws.append(hdr)
for c in ws[1]: c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="FFFF00")
for r in range(2, 193):
    ws.cell(r, 6).value = f'=IF(LEFT(C{r},2)="優良","加分",IF(LEFT(C{r},2)="不良","扣分",""))'
    ws.cell(r, 7).number_format = "0.0"
items = wb.create_sheet("加扣分項目")
cats = {
  "優良事蹟(一)": ["24.對指派或推選擔任之工作，執行得力，成績優良者", "25.公差勤務工作努力，績效優良者"],
  "優良事蹟(二)": ["30.一週內務經常保持整潔者", "31.在校精神飽滿服裝整潔足資表率者"],
  "不良事蹟(一)": ["22.對指派或推選擔任之工作執行錯誤，情節尚輕者", "63.內務不整者"],
  "不良事蹟(二)": ["59.儀容不整，指甲未修者", "69.其它違反生活規範規定，情節輕微者"],
}
items["A1"] = "說明文字"
for ci, (cat, its) in enumerate(cats.items(), 1):
    items.cell(14, ci).value = cat
    for ri, it in enumerate(its, 15): items.cell(ri, ci).value = it
dv = DataValidation(type="list", formula1="'加扣分項目'!$A$14:$D$14"); ws.add_data_validation(dv); dv.add("C2:C192")
wb.save(FX + "points_template.xlsx")

# 獎懲範本
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "工作表1"
ws.append(["學號", "姓名", "獎懲依據(條)", "獎懲依據(款)", "獎懲依據(目)", "獎懲事由", "額度", "加重減輕事由 (若無可不選擇)", "加重減輕額度"])
rules = wb.create_sheet("獎懲項目")
for row in [
  ["", "", "", "", "", "", "第19條", "申誡一次", "", "加重減輕事由"],
  ["第19條", "第1款", "", "", "測試事由甲。", "", 1, "申誡一次", "", "行為之動機"],
  ["第20條", "第17款", "第1目", "", "測試事由乙。", "", "第20條", "小過一次"],
  ["第20條", "第17款", "第2目", "", "測試事由丙。", "", 2, "小過一次"],
  ["", "", "", "", "", "", 13, "以下空白"]]:
    rules.append(row)
wb.save(FX + "reward_template.xlsx")
print("ok")
